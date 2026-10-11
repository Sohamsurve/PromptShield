"""Build the hackathon pitch deck as PowerPoint and PDF. Same eight slides."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from reportlab.lib.colors import Color, white
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent
LOGO = OUT.parents[1] / "apps" / "web" / "public" / "assets" / "favicon-512.png"
W_IN, H_IN = 13.333, 7.5
W, H = W_IN * 72, H_IN * 72

INK = (0x0E, 0x17, 0x26)
PAPER = (0xF3, 0xF5, 0xF9)
SURFACE = (0xFF, 0xFF, 0xFF)
SLATE = (0x5B, 0x66, 0x7A)
UV = (0x6A, 0x4D, 0xFF)
ALLOW = (0x13, 0x8A, 0x5B)
SAN = (0xB2, 0x6B, 0x00)
QUA = (0x0B, 0x7B, 0xA7)
BLOCK = (0xC6, 0x28, 0x39)
ESC = (0xA4, 0x23, 0x7A)


def rgb(t):
    return RGBColor(*t)


def pdf_color(t):
    return Color(t[0] / 255, t[1] / 255, t[2] / 255)


def _no_autofit(shape):
    body = shape.text_frame._txBody
    pr = body.find(qn("a:bodyPr"))
    if pr is not None:
        for child in list(pr):
            if child.tag == qn("a:spAutoFit"):
                pr.remove(child)
        pr.set("anchor", "t")


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W_IN)
        self.prs.slide_height = Inches(H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.pdf = canvas.Canvas(str(OUT / "PromptShield_pitch.pdf"), pagesize=(W, H))
        self.slide = None
        self._logo_file = type("P", (), {"name": str(LOGO)})()

    def new(self):
        self.slide = self.prs.slides.add_slide(self.blank)
        bg = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W_IN), Inches(H_IN))
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb(PAPER)
        bg.line.fill.background()
        self.pdf.setFillColor(pdf_color(PAPER))
        self.pdf.rect(0, 0, W, H, fill=1, stroke=0)

    def rect(self, x, y, w, h, fill, line=INK, lw=1.5, shadow=False):
        """x, y, w, h in inches from the top-left."""
        if shadow:
            sh = self.slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(x + 0.06), Inches(y + 0.06), Inches(w), Inches(h)
            )
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(INK)
            sh.line.fill.background()
            self.pdf.setFillColor(pdf_color(INK))
            self.pdf.rect((x + 0.06) * 72, H - (y + 0.06 + h) * 72, w * 72, h * 72, fill=1, stroke=0)
        box = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
        box.line.color.rgb = rgb(line)
        box.line.width = Pt(lw)
        self.pdf.setStrokeColor(pdf_color(line))
        self.pdf.setFillColor(pdf_color(fill))
        self.pdf.setLineWidth(lw)
        self.pdf.rect(x * 72, H - (y + h) * 72, w * 72, h * 72, fill=1, stroke=1)
        return box

    def text(self, value, x, y, w, h, size=18, color=INK, bold=False, align="left"):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        _no_autofit(box)
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        run = p.add_run()
        run.text = value
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        self.pdf.setFillColor(pdf_color(color))
        self.pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        # ReportLab draws from the baseline. Approximate a top-aligned block.
        leading = size * 1.25
        font = "Helvetica-Bold" if bold else "Helvetica"
        lines = []
        max_w = w * 72 - 4
        for paragraph in value.split("\n"):
            cur = ""
            for word in paragraph.split():
                trial = word if not cur else cur + " " + word
                if self.pdf.stringWidth(trial, font, size) <= max_w:
                    cur = trial
                else:
                    if cur:
                        lines.append(cur)
                    cur = word
            lines.append(cur)
        if not lines:
            lines = [""]
        yy = H - y * 72 - size
        for line in lines:
            if align == "center":
                self.pdf.drawCentredString((x + w / 2) * 72, yy, line)
            else:
                self.pdf.drawString(x * 72, yy, line)
            yy -= leading

    def logo(self, x, y, size):
        self.slide.shapes.add_picture(self._logo_file.name, Inches(x), Inches(y), Inches(size), Inches(size))
        self.pdf.drawImage(
            self._logo_file.name,
            x * 72,
            H - (y + size) * 72,
            size * 72,
            size * 72,
            mask="auto",
            preserveAspectRatio=True,
            anchor="c",
        )

    def footer(self, n, total=8):
        self.logo(0.42, 7.05, 0.32)
        self.text("PromptShield  ·  ET AI Hackathon  ·  Problem 2", 0.84, 7.12, 8, 0.28, 11, SLATE)
        self.text(f"{n}  /  {total}", 11.4, 7.12, 1.5, 0.28, 11, SLATE, align="center")

    def finish(self):
        self.pdf.showPage()

    def save(self):
        self.prs.save(OUT / "PromptShield_pitch.pptx")
        self.pdf.save()


def build():
    d = Deck()

    d.new()
    d.rect(0, 0, 0.18, H_IN, UV, UV, 0)
    d.logo(0.55, 1.72, 1.15)
    d.text("ET AI HACKATHON  ·  AGENTIC EDITION  ·  PROBLEM 2", 1.95, 1.55, 10, 0.35, 14, UV, True)
    d.text("PromptShield", 1.95, 1.9, 10, 0.9, 54, INK, True)
    d.text("A prompt-injection firewall for AI agents.", 0.55, 2.85, 10, 0.5, 24, INK)
    d.text(
        "It checks untrusted text and files before the model sees them,\nneutralises the hidden instruction, and still lets the ordinary text through.",
        0.55,
        3.5,
        10,
        0.8,
        16,
        SLATE,
    )
    d.rect(0.55, 4.7, 7.6, 0.85, SURFACE, shadow=True)
    d.text("GitHub", 0.75, 4.82, 1.3, 0.3, 12, SLATE, True)
    d.text("https://github.com/Sohamsurve/PromptShield", 0.75, 5.1, 7.1, 0.35, 16, INK, True)
    d.finish()

    d.new()
    d.text("The problem", 0.5, 0.32, 8, 0.5, 32, INK, True)
    d.text("Agents treat retrieved content as instructions.", 0.5, 0.95, 12, 0.4, 18, SLATE)
    cards = [
        ("01", "Hidden instruction", "A line in an email, PDF, or web page can override the task the user actually asked for."),
        ("02", "Side effect", "The same line can ask the agent to reveal a secret or call a tool such as send-email."),
        ("03", "Wrong tool", "A toxicity filter does not catch this. Tool permissions still matter, and this firewall does not replace them."),
    ]
    for i, (num, title, body) in enumerate(cards):
        x = 0.5 + i * 4.2
        d.rect(x, 1.7, 3.95, 3.5, SURFACE, shadow=True)
        d.text(num, x + 0.25, 1.95, 3.4, 0.45, 22, UV, True)
        d.text(title, x + 0.25, 2.55, 3.4, 0.7, 22, INK, True)
        d.text(body, x + 0.25, 3.4, 3.45, 1.5, 15, SLATE)
    d.footer(2)
    d.finish()

    d.new()
    d.text("What the prototype does", 0.5, 0.32, 12, 0.5, 32, INK, True)
    rows = [
        ("Scan first", "Every prompt and file is checked before it reaches the agent."),
        ("Redact the span", "Only the offending sentence is removed. The rest of the document stays."),
        ("Safe alternative", "On block or quarantine, the caller gets the benign remainder, not the attack."),
        ("Tool firewall", "A tainted send-email is denied even when text scanning is switched off."),
    ]
    for i, (title, body) in enumerate(rows):
        y = 1.15 + i * 1.25
        d.rect(0.5, y, 8.3, 1.1, SURFACE, shadow=True)
        d.text(title, 0.75, y + 0.15, 7.8, 0.35, 18, INK, True)
        d.text(body, 0.75, y + 0.52, 7.8, 0.4, 14, SLATE)
    verdicts = [("Allowed", ALLOW), ("Sanitised", SAN), ("Quarantined", QUA), ("Blocked", BLOCK), ("Needs review", ESC)]
    d.text("Five answers", 9.1, 1.15, 3.6, 0.35, 16, INK, True)
    for i, (word, color) in enumerate(verdicts):
        y = 1.65 + i * 0.85
        d.rect(9.1, y, 3.7, 0.72, SURFACE, shadow=True)
        d.rect(9.1, y, 0.12, 0.72, color, color, 0)
        d.text(word, 9.4, y + 0.18, 3.2, 0.4, 16, color, True)
    d.footer(3)
    d.finish()

    d.new()
    d.text("How one scan works", 0.5, 0.32, 12, 0.5, 32, INK, True)
    steps = [
        ("1", "Extract", "Text, HTML, PDF, Word, email, Markdown, JSON, source, and images. Hidden text is surfaced."),
        ("2", "Detect", "Signatures first, then a small classifier. Encoded text is decoded, with a depth cap."),
        ("3", "Decide", "Fusion picks the verdict. The language-model judge never decides alone."),
        ("4", "Neutralise", "The span is redacted in the original text. A second scan checks the result."),
        ("5", "Guard tools", "Before a side effect, the tool firewall checks taint, destination, and canaries."),
    ]
    for i, (n, title, body) in enumerate(steps):
        y = 1.15 + i * 1.05
        d.rect(0.5, y, 12.3, 0.92, SURFACE, shadow=True)
        d.text(n, 0.7, y + 0.22, 0.5, 0.45, 22, UV, True)
        d.text(title, 1.35, y + 0.12, 2.4, 0.65, 18, INK, True)
        d.text(body, 3.9, y + 0.25, 8.6, 0.5, 15, SLATE)
    d.footer(4)
    d.finish()

    d.new()
    d.text("What the video shows", 0.5, 0.32, 12, 0.5, 32, INK, True)
    beats = [
        ("Landing", "One specimen email. Ultraviolet marks only the hidden instruction."),
        ("Arena", "The same email, firewall off and on. The unprotected agent follows the hidden line."),
        ("X-ray", "Revealed view, reason codes, and the safe alternative without the canary."),
        ("Clean pass", "A normal security article stays Allowed."),
        ("Last line", "Text scanning off. A tainted send-email is still denied."),
        ("Accounts", "A caller scans a file. An admin sees each account’s usage."),
    ]
    for i, (title, body) in enumerate(beats):
        col = i % 3
        row = i // 3
        x = 0.5 + col * 4.2
        y = 1.25 + row * 2.55
        d.rect(x, y, 3.95, 2.3, SURFACE, shadow=True)
        d.text(f"0{i + 1}", x + 0.25, y + 0.25, 3.4, 0.4, 16, UV, True)
        d.text(title, x + 0.25, y + 0.7, 3.45, 0.45, 20, INK, True)
        d.text(body, x + 0.25, y + 1.25, 3.45, 0.8, 14, SLATE)
    d.footer(5)
    d.finish()

    d.new()
    d.text("Measured on the frozen golden suite", 0.5, 0.28, 12, 0.45, 28, INK, True)
    d.text("Judge off. Shipped policy. 11 October 2026. Not tuned on these rows.", 0.5, 0.8, 12, 0.3, 14, SLATE)
    stats = [
        ("0.61", "Acted recall", "900 attacks"),
        ("0.48", "Correct type", "900 attacks"),
        ("0 / 500", "False positives", "Ordinary documents"),
        ("26 ms", "Latency p95", "Without the judge"),
    ]
    for i, (num, label, note) in enumerate(stats):
        x = 0.5 + i * 3.2
        d.rect(x, 1.35, 3.0, 2.15, SURFACE, shadow=True)
        d.text(num, x + 0.15, 1.55, 2.7, 0.7, 32, INK, True, "center")
        d.text(label, x + 0.15, 2.35, 2.7, 0.4, 14, INK, True, "center")
        d.text(note, x + 0.15, 2.8, 2.7, 0.4, 12, SLATE, False, "center")
    notes = [
        ("Caught in full", "Indirect injection, 100 of 100 on this suite."),
        ("Still weak", "Multi-step jailbreaks: acted 0.35, correct type 0.00."),
        ("Other pack", "Imported sentences: acted 0.18, false positives 0 of 32."),
    ]
    for i, (title, body) in enumerate(notes):
        x = 0.5 + i * 4.2
        d.rect(x, 3.8, 3.95, 1.7, SURFACE, shadow=True)
        d.text(title, x + 0.25, 4.0, 3.45, 0.4, 16, INK, True)
        d.text(body, x + 0.25, 4.5, 3.45, 0.75, 14, SLATE)
    d.text("Full table: docs/CLAIMS.md  ·  Report: eval/golden/reports/20261011T090211Z", 0.5, 5.7, 12, 0.35, 13, SLATE)
    d.footer(6)
    d.finish()

    d.new()
    d.text("What we declare", 0.5, 0.32, 12, 0.5, 32, INK, True)
    d.rect(0.5, 1.15, 6.0, 2.3, SURFACE, shadow=True)
    d.text("F1  /  D1", 0.75, 1.4, 5.5, 0.6, 32, UV, True)
    d.text("Provisional. Nine attack types and the listed formats are in the product. An independent sealed run has not been scored.", 0.75, 2.15, 5.5, 1.0, 15, SLATE)
    d.rect(6.8, 1.15, 6.0, 2.3, SURFACE, shadow=True)
    d.text("Not claimed", 7.05, 1.4, 5.5, 0.45, 22, BLOCK, True)
    d.text("F3 and D3. Known-pattern recall of 0.95. Held-out recall of 0.85. A length-matched false-positive rate at or under 2 percent.", 7.05, 2.05, 5.5, 1.1, 15, SLATE)
    limits = [
        "Exact white-on-white text is not in the pixels, so OCR cannot read it.",
        "The Arena agent is scripted. A live language-model agent was not the measured subject.",
        "The judge can add a finding. It cannot block a document by itself.",
    ]
    for i, line in enumerate(limits):
        y = 3.7 + i * 0.9
        d.rect(0.5, y, 12.3, 0.78, SURFACE, shadow=True)
        d.text(line, 0.75, y + 0.2, 11.8, 0.45, 15, INK)
    d.footer(7)
    d.finish()

    d.new()
    d.rect(0, 0, 0.18, H_IN, UV, UV, 0)
    d.logo(0.55, 2.05, 1.15)
    d.text("PromptShield", 1.95, 2.2, 10, 0.9, 48, INK, True)
    d.text("Check the content. Remove the attack. Keep the document.", 0.55, 3.2, 11, 0.5, 22, INK)
    d.rect(0.55, 4.15, 7.6, 0.85, SURFACE, shadow=True)
    d.text("GitHub", 0.75, 4.27, 1.3, 0.28, 12, SLATE, True)
    d.text("https://github.com/Sohamsurve/PromptShield", 0.75, 4.55, 7.1, 0.35, 16, INK, True)
    d.footer(8)
    d.finish()

    d.save()


if __name__ == "__main__":
    build()
    print("wrote", OUT / "PromptShield_pitch.pptx")
    print("wrote", OUT / "PromptShield_pitch.pdf")
